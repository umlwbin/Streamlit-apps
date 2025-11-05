import streamlit as st
import zipfile
from zipfile import ZipFile
import os
import io
import sys

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from io import BytesIO
import pandas as pd

#Output Path
path=os.path.abspath(os.curdir)

#Add Modules
sys.path.append(f'{path}/Modules')

import  session_initializer

def download_output():
    def on_download():
        st.balloons()
        #Set Next button states to False
        session_initializer.reset_widget_flags() #Reset all the next buttons that are not directly widgetkeys.

    st.markdown('## 📦 Download Processed Files')
    if len(st.session_state.current_data) > 0:
        st.markdown('### 📑 CSV')

    #Check if this was a Merge Function
    merge_check=[s for s in st.session_state.current_data.keys() if 'merge' in s]

    #Merge Task
    if merge_check:
        filename=merge_check[0] # 'merged_output.csv'
        df=st.session_state.current_data['merged_output.csv']
        st.download_button(
            label=f"⬇️ Download {filename}",
            data=df.to_csv(index=False).encode('utf-8'),
            file_name=f"{filename}",
            mime="text/csv",
            on_click=on_download,
            icon=":material/download:"
        )

    # All Other Tasks
    else:

        if len(st.session_state.current_data) == 1:
            filename, df = next(iter(st.session_state.current_data.items()))
            st.write(filename)
            st.download_button(
                label=f"Download {filename}",
                data=df.to_csv(index=False).encode('utf-8'),
                file_name=f"{filename[:-4]}_cleaned.csv",
                mime="text/csv",
                on_click=on_download,
                icon=":material/download:"
            )

        if len(st.session_state.current_data) > 1:
            dfs_dict=st.session_state.current_data
            zip_buffer = io.BytesIO()
            with ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
                for filename, df in dfs_dict.items():
                    csv_bytes = df.to_csv(index=False).encode('utf-8')
                    zip_file.writestr(f"{filename[:-4]}_cleaned.csv", csv_bytes)
            zip_buffer.seek(0)

            zip_data=zip_buffer

            st.download_button(
                label=f"Download All as ZIP",
                data=zip_data,
                file_name="output_data.zip",
                mime="application/zip",
                on_click=on_download,
                icon=":material/download:"
            )

# EXCEL-------------------------------------------------------------------------------------------------
def to_excel_with_formatting(df, freeze_header=False):

    output = BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = "Cleaned Data"

    # Identify datetime columns
    datetime_cols = [col for col in df.columns if pd.api.types.is_datetime64_any_dtype(df[col])]

    # Styles
    header_font = Font(bold=True)
    header_fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")
    datetime_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")  # soft yellow

    # Write headers
    for col_num, column_title in enumerate(df.columns, 1):
        cell = ws.cell(row=1, column=col_num, value=column_title)
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        cell.fill = header_fill

    # Write data rows
    for row_num, row_data in enumerate(df.values, 2):
        for col_num, cell_value in enumerate(row_data, 1):
            if pd.isna(cell_value): # Convert <NA> to None
                cell_value = None
            cell = ws.cell(row=row_num, column=col_num, value=cell_value)
            if df.columns[col_num - 1] in datetime_cols:
                cell.fill = datetime_fill

    # Auto-adjust column widths
    for column_cells in ws.columns:
        length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
        ws.column_dimensions[column_cells[0].column_letter].width = length + 2

    # Freeze header row if requested
    if freeze_header:
        ws.freeze_panes = "A2"

    wb.save(output)
    output.seek(0)
    return output

def excel_download():
    if len(st.session_state.current_data) == 1:
        filename, df = next(iter(st.session_state.current_data.items()))
        st.markdown(" ")
        st.markdown('### 📑 EXCEL')
        freeze = st.checkbox("Freeze header row in Excel")

        excel_data = to_excel_with_formatting(df, freeze_header=freeze)
        st.download_button(
            label="Download Excel File with Formatting",
            data=excel_data,
            file_name=filename[:-4]+"_cleaned.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            icon=":material/download:"
        )
